"""
Utility helpers for processing file attachments uploaded via the Streamlit UI.

Supported attachment types:
  - Images (PNG / JPG / JPEG / WEBP / GIF) -> base64 data URI for vision LLMs
  - PDF                                    -> plain text via pypdf
  - DOCX                                   -> paragraphs, tables and equation text via python-docx
  - PPTX                                   -> slide text and tables via python-pptx
  - XLSX                                   -> cell values via openpyxl
  - HTML / HTM                             -> visible text (scripts / styles dropped)
  - TXT / MD / CSV / JSON / TEX            -> raw decoded text

Legacy binary formats (.doc, .ppt, .xls) are not supported; save them as the
modern Office format first.
"""
from __future__ import annotations

import base64
import io
import logging
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import Callable, Dict, Optional, Tuple

logger = logging.getLogger(__name__)

# -- Limits -----------------------------------------------------------------

MAX_FILE_BYTES = 10 * 1024 * 1024            # per-file upload cap
MAX_FILES = 5                                # attachments per message
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024   # zip-bomb guard for Office files
TOTAL_TEXT_BUDGET = 12000                    # characters shared by all text attachments

# -- Supported types --------------------------------------------------------

_MIME_MAP = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
             "webp": "image/webp", "gif": "image/gif"}

SUPPORTED_IMAGE_EXTS = set(_MIME_MAP.keys())
SUPPORTED_DOC_EXTS   = {"pdf", "docx", "pptx", "xlsx", "html", "htm",
                        "txt", "md", "csv", "json", "tex"}
SUPPORTED_EXTS       = SUPPORTED_IMAGE_EXTS | SUPPORTED_DOC_EXTS


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _truncate(text: str, max_chars: int) -> str:
    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[truncated at {max_chars} chars]"
    return text


def _check_zip_size(file_bytes: bytes) -> None:
    """Reject Office files whose declared uncompressed size is unreasonable."""
    with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("Archive expands to an unreasonable size")


# -- Extractors -------------------------------------------------------------

def extract_text_from_pdf(file_bytes: bytes, max_chars: int = 6000) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    parts: list[str] = []
    total = 0
    for page in reader.pages:
        t = (page.extract_text() or "").strip()
        if t:
            parts.append(t)
            total += len(t)
        if total > max_chars:
            break
    return _truncate("\n\n".join(parts), max_chars)


def extract_text_from_txt(file_bytes: bytes, max_chars: int = 6000) -> str:
    return _truncate(file_bytes.decode("utf-8", errors="replace"), max_chars)


def extract_text_from_docx(file_bytes: bytes, max_chars: int = 6000) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip_size(file_bytes)
    W_T = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    M_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"

    doc = Document(io.BytesIO(file_bytes))
    lines: list[str] = []
    total = 0
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            # w:t = normal text, m:t = Office Math (equation) text
            text = "".join(el.text or "" for el in block._p.iter() if el.tag in (W_T, M_T))
            line = text.strip()
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                rows.append(" | ".join(cells))
            line = "\n".join(rows)
        else:
            continue
        if line:
            lines.append(line)
            total += len(line)
        if total > max_chars:
            break
    return _truncate("\n".join(lines), max_chars)


def extract_text_from_pptx(file_bytes: bytes, max_chars: int = 6000) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    _check_zip_size(file_bytes)

    def shape_text(shape) -> list[str]:
        out: list[str] = []
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in shape.shapes:
                out.extend(shape_text(sub))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            if shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                out.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
        return out

    prs = Presentation(io.BytesIO(file_bytes))
    lines: list[str] = []
    total = 0
    for n, slide in enumerate(prs.slides, start=1):
        lines.append(f"[Slide {n}]")
        for shape in slide.shapes:
            for t in shape_text(shape):
                lines.append(t)
                total += len(t)
        if total > max_chars:
            break
    return _truncate("\n".join(lines), max_chars)


def extract_text_from_xlsx(file_bytes: bytes, max_chars: int = 6000) -> str:
    from openpyxl import load_workbook

    _check_zip_size(file_bytes)
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        lines: list[str] = []
        total = 0
        for ws in wb.worksheets:
            lines.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if not any(cells):
                    continue
                line = "\t".join(cells).rstrip()
                lines.append(line)
                total += len(line)
                if total > max_chars:
                    break
            if total > max_chars:
                break
    finally:
        wb.close()
    return _truncate("\n".join(lines), max_chars)


class _HTMLTextExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "head", "template"}
    _BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
              "table", "section", "article", "ul", "ol", "pre", "blockquote"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def extract_text_from_html(file_bytes: bytes, max_chars: int = 6000) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(file_bytes.decode("utf-8", errors="replace"))
    parser.close()
    text = "".join(parser.parts)
    lines = [" ".join(ln.split()) for ln in text.splitlines()]
    return _truncate("\n".join(ln for ln in lines if ln), max_chars)


_EXTRACTORS: Dict[str, Callable[[bytes, int], str]] = {
    "pdf": extract_text_from_pdf,
    "docx": extract_text_from_docx,
    "pptx": extract_text_from_pptx,
    "xlsx": extract_text_from_xlsx,
    "html": extract_text_from_html,
    "htm": extract_text_from_html,
}


def extract_text(file_bytes: bytes, filename: str, max_chars: int = 6000) -> str:
    """Dispatch to the right extractor; unknown text-like extensions are decoded as UTF-8."""
    return _EXTRACTORS.get(_ext(filename), extract_text_from_txt)(file_bytes, max_chars)


# -- Image -> base64 data URI -----------------------------------------------

def image_to_base64_uri(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    mime = _MIME_MAP.get(_ext(filename), "image/png")
    b64  = base64.b64encode(file_bytes).decode("ascii")
    return f"data:{mime};base64,{b64}", mime


def is_image(filename: str) -> bool:
    return _ext(filename) in SUPPORTED_IMAGE_EXTS


# -- Attachment container ---------------------------------------------------

@dataclass
class Attachment:
    name: str
    is_image: bool
    preview: object                      # raw bytes for images, text otherwise
    image_uri: Optional[str] = None
    text: Optional[str] = None


def build_attachment(name: str, file_bytes: bytes, max_chars: int = 6000) -> Attachment:
    """
    Turn an uploaded file into an Attachment.
    Raises ValueError for unsupported / oversized files; other exceptions mean
    the file could not be parsed (callers should show a generic message).
    """
    if _ext(name) not in SUPPORTED_EXTS:
        raise ValueError("Unsupported file type")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds {MAX_FILE_BYTES // (1024 * 1024)} MB limit")

    if is_image(name):
        uri, _ = image_to_base64_uri(file_bytes, name)
        return Attachment(name=name, is_image=True, preview=file_bytes, image_uri=uri)

    try:
        text = extract_text(file_bytes, name, max_chars)
    except Exception:
        logger.exception("Failed to extract text from attachment")
        raise
    return Attachment(name=name, is_image=False, preview=text, text=text)
